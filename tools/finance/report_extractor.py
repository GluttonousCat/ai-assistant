# -*- encoding: utf-8 -*-
"""
研报文本提取器
将本地研报文件 (PDF/DOCX/TXT/MD) 提取为纯文本, 供研报校验/分析链路使用.

策略:
- PDF: 优先 PyMuPDF(fitz), 缺失时尝试 pdfplumber / pypdf
- DOCX: python-docx, 缺失时退化为 zipfile 直接解 word/document.xml
- TXT/MD: utf-8 / gbk 读入
"""
from __future__ import annotations


import os
import re
import zipfile
from typing import Dict, Optional
from xml.etree import ElementTree

from core.logger import get_logger
logger = get_logger(__name__)

SUPPORTED_EXTS = {".pdf", ".docx", ".doc", ".txt", ".md", ".markdown"}


class ReportExtractor:
    """研报文件 → 纯文本"""

    # 常见噪音行 (页眉页脚/页码) 过滤
    NOISE_PATTERNS = [
        re.compile(r"^\s*[-—–]\s*\d+\s*[-—–]\s*$"),          # 页码 - 3 -
        re.compile(r"^\s*\d+\s*/\s*\d+\s*$"),                   # 3/12
        re.compile(r"^\s*仅供内部使用.*$", re.I),
        re.compile(r"^\s*免责声明[:：]?.*$"),
        re.compile(r"^\s*\[\s*Table_PageBreak\s*\]\s*$"),
    ]

    def extract(self, file_path: str) -> Optional[str]:
        """提取文件文本, 失败返回 None"""
        if not os.path.exists(file_path):
            logger.warning("文件不存在: %s", file_path)
            return None

        ext = os.path.splitext(file_path)[1].lower()
        try:
            if ext == ".pdf":
                text = self._extract_pdf(file_path)
            elif ext in (".docx", ".doc"):
                text = self._extract_docx(file_path)
            elif ext in (".txt", ".md", ".markdown"):
                text = self._extract_text(file_path)
            else:
                logger.warning("不支持的格式: %s", ext)
                return None
        except Exception as e:
            logger.error("提取失败 %s: %s", file_path, e)
            return None

        if not text:
            return None
        return self.clean(text)

    # ---------- 各格式实现 ----------
    def _extract_pdf(self, path: str) -> str:
        # PDF 原语统一走 tools/pdf (fitz→pdfplumber→pypdf 降级链)
        from tools.pdf import extract_text
        return extract_text(path)

    def _extract_docx(self, path: str) -> str:
        # .doc (老格式) 为二进制, zipfile 解不出时抛出, 由调用方标记失败
        try:
            import docx  # python-docx

            document = docx.Document(path)
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except ImportError:
            pass

        # 退化: zipfile 直接读 document.xml
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml")
        root = ElementTree.fromstring(xml)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        paragraphs = []
        for p in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
            texts = [t.text or "" for t in p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")]
            paragraphs.append("".join(texts))
        return "\n".join(paragraphs)

    def _extract_text(self, path: str) -> str:
        for enc in ("utf-8", "gbk", "gb18030", "latin-1"):
            try:
                with open(path, "r", encoding=enc) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        return ""

    # ---------- 清洗 ----------
    def clean(self, text: str) -> str:
        """噪音过滤 + 压缩空白"""
        lines = []
        for line in text.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if any(p.match(stripped) for p in self.NOISE_PATTERNS):
                continue
            lines.append(stripped)
        return "\n".join(lines)

    def extract_meta(self, file_path: str) -> Dict:
        """提取基础元信息 (无文件读取, 纯命名统计)"""
        name = os.path.basename(file_path)
        return {
            "file_name": name,
            "ext": os.path.splitext(name)[1].lower(),
            "size": os.path.getsize(file_path) if os.path.exists(file_path) else 0,
        }


# ---------- 便捷单例 ----------
_extractor: Optional[ReportExtractor] = None


def get_extractor() -> ReportExtractor:
    global _extractor
    if _extractor is None:
        _extractor = ReportExtractor()
    return _extractor